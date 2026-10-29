"""Approved delivery grades preserve verified work without inventing unread evidence."""
import copy
import io
import json

from docx import Document
from fastapi.testclient import TestClient
import pytest

from engines.common.llm_client import LLMContentRejected
from engines.insight_engine import main as insight
from engines.report_engine import main as report


DOC = {"id": "d1", "title": "续航实测", "content": "实测续航为八小时。", "source_type": "news",
       "source_name": "测试新闻", "url": "https://example.com/a", "published_at": "2026-10-29"}
DIMENSIONS = ["background", "heat", "sentiment", "group", "deep_cause"]


def verified_request():
    span = {"document_id": "d1", "start": 0, "end": 9, "material_class": "direct"}
    return {
        "resilient": True, "delivery_policy": "best_available", "delivery_level": "full", "api_key": "offline-key",
        "title": "续航核验报告", "documents": [copy.deepcopy(DOC)], "summary": "已核验的摘要：实测八小时。",
        "sentiments": [{"document_id": "d1", "sentiment": "neutral", "score": 0.5}],
        "topics": [{"id": "t1", "name": "续航反馈", "keywords": ["续航"], "doc_ids": ["d1"], "doc_count": 1}],
        "dimensions": [{"id": dim, "name": dim, "findings": f"核验结论-{dim}", "data_points": ["1 篇直接材料"],
                        "deep_read": f"核验解读-{dim}", "trend": f"材料不足以判断趋势-{dim}",
                        "quotes": [{"document_id": "d1", "text": "实测续航为八小时。", "source": "测试新闻"}],
                        "quote_locations": [{**span, "quote_text": "实测续航为八小时。"}]} for dim in DIMENSIONS],
        "insight_material_coverage": {"version": "complete-chunks-v1", "phases": [
            {"phase": phase, "read_spans": [copy.deepcopy(span)], "unprocessed": []}
            for phase in ["classification", *["dimension:" + dim for dim in DIMENSIONS]]
        ] + [{"phase": "summary", "read_spans": [], "unprocessed": [], "input_basis": "verified_dimension_findings"}]},
    }


class PartialProvider:
    async def chat_json(self, model, messages, **kwargs):
        prompt = messages[-1]["content"]
        if "【分析维度】" in prompt:
            if "id=heat" in prompt and "原文片段[8000:" in prompt:
                raise LLMContentRejected()
            return {"findings": "本批已完成核验的结论", "data_points": ["本批1个片段"], "quotes": [],
                    "deep_read": "本批材料支持的解释", "trend": "仅限已读范围，趋势未知"}
        if "sentiments" in prompt:
            return {"sentiments": [{"document_id": "d1", "sentiment": "neutral", "score": 0.5,
                                     "level": "中性", "confidence": 0.9, "emotions": {}}],
                    "topics": [{"id": "t1", "name": "续航", "keywords": ["续航"], "doc_ids": ["d1"], "doc_count": 1}]}
        pytest.fail("A refused material batch must not be repurchased through summary writing")


def test_resilient_insight_retains_verified_partial_dimension_and_exact_unread_tail(monkeypatch):
    monkeypatch.setattr(insight, "build_client", lambda *_args, **_kwargs: PartialProvider())
    doc = {**DOC, "content": "续航实测。" * 1800}
    response = TestClient(insight.app).post("/analyze", json={"resilient": True, "api_key": "offline-key", "documents": [doc]})
    assert response.status_code == 200
    data = response.json()
    assert data["delivery_level"] == "limited"
    assert len(data["dimensions"]) == 5
    heat = next(d for d in data["dimensions"] if d["id"] == "heat")
    assert heat["findings"] and heat["complete"] is False
    phase = next(p for p in data["material_coverage"]["phases"] if p["phase"] == "dimension:heat")
    assert [(s["start"], s["end"]) for s in phase["read_spans"]] == [(0, 8000)]
    assert [(s["start"], s["end"], s["code"]) for s in phase["unprocessed"]] == [(8000, 9000, "llm_content_rejected")]
    assert data["summary"] and data["delivery_reasons"]


def test_resilient_insight_without_model_preserves_material_receipt(monkeypatch):
    monkeypatch.setattr(insight, "LLM_API_KEY", "")
    monkeypatch.setattr(insight, "build_client", lambda *_args, **_kwargs: pytest.fail("brief cannot buy model calls"))
    response = TestClient(insight.app).post("/analyze", json={"resilient": True, "documents": [DOC]})
    assert response.status_code == 200
    data = response.json()
    assert data["delivery_level"] == "brief" and data["delivery_reasons"]
    assert data["sentiments"] == [] and data["dimensions"] == []
    assert data["material_coverage"]["phases"][0]["read_spans"] == []


@pytest.mark.parametrize("grade", ["brief", "empty"])
def test_explicit_receipts_need_no_model_and_keep_html_word_snapshot(monkeypatch, grade):
    monkeypatch.setattr(report, "build_client", lambda *_args, **_kwargs: pytest.fail("receipt must not call a model"))
    raw = {**DOC, "id": "raw", "title": "未完成相关性核验 <原始材料>", "semantic_classification": "uncertain"}
    body = {"resilient": True, "delivery_level": grade, "delivery_reasons": ["模型不可用"], "api_key": "present-but-forbidden",
            "title": "资料交付", "unverified_documents": [raw] if grade == "brief" else [],
            "sentiments": [{"document_id": "raw", "sentiment": "positive"}],
            "dimensions": [{"id": "heat", "findings": "禁止展示的语义判断"}]}
    response = TestClient(report.app).post("/generate", json=body)
    assert response.status_code == 200
    data = response.json()
    assert data["delivery_level"] == grade and data["delivery_reasons"] == ["模型不可用"]
    assert "资料简报" in data["content"] if grade == "brief" else "空结果回执" in data["content"]
    assert "禁止展示的语义判断" not in data["content"] and "正面" not in data["content"]
    if grade == "brief":
        assert "未完成相关性核验 &lt;原始材料&gt;" in data["content"]
        assert DOC["url"] in data["content"] and "2026-10-29" in data["content"]
    saved = report.render_saved_docx(data["content"])
    word = Document(io.BytesIO(saved))
    text = "\n".join(p.text for p in word.paragraphs) + "\n" + "\n".join(c.text for t in word.tables for row in t.rows for c in row.cells)
    assert "模型不可用" in text and "禁止展示的语义判断" not in text


def test_writer_failure_compiles_every_verified_component_and_citation_without_downgrade(monkeypatch):
    class BrokenWriter:
        async def chat_json(self, *_args, **_kwargs):
            raise LLMContentRejected()
    monkeypatch.setattr(report, "build_client", lambda *_args, **_kwargs: BrokenWriter())
    response = TestClient(report.app).post("/generate", json=verified_request())
    assert response.status_code == 200
    data = response.json()
    assert data["delivery_level"] == "full"
    assert "已核验的摘要：实测八小时。" in data["content"]
    for dim in DIMENSIONS:
        for value in (f"核验结论-{dim}", f"核验解读-{dim}", f"材料不足以判断趋势-{dim}"):
            assert value in data["content"]
    assert "实测续航为八小时。" in data["content"] and "[0:9]" in data["content"]
    phase = data["material_coverage"]["phases"][0]
    assert phase["input_basis"] == "verified_components"
    assert phase["read_spans"] == [] and not phase["unprocessed"]


def test_unsafe_component_cannot_be_upgraded_to_full_by_deterministic_writer(monkeypatch):
    body = verified_request()
    body["api_key"] = ""
    body["dimensions"][0]["quotes"][0]["text"] = "原文没有的伪造引语"
    body["insight_material_coverage"]["phases"][1]["unprocessed"] = [
        {"document_id": "d1", "start": 0, "end": 9, "reason": "供应商拒绝"}]
    monkeypatch.setattr(report, "LLM_API_KEY", "")
    response = TestClient(report.app).post("/generate", json=body)
    assert response.status_code == 200
    data = response.json()
    assert data["delivery_level"] == "limited" and data["delivery_reasons"]
    assert "原文没有的伪造引语" not in data["content"]
    assert "情感分类覆盖 1/1" in data["content"]
