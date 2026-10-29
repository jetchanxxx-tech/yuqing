"""Saved report export preserves published text without buying a new model result."""
import io

from docx import Document
from fastapi.testclient import TestClient
import pytest

from engines.report_engine import main as report


HTML = '''<!doctype html><html><head><style>hidden-css</style><title>hidden-title</title></head>
<body><h1>盘古舆情·第一版</h1><p>摘要保持“原文”。<br>第二行</p>
<table><tr><th>文档编号</th><th>来源／日期</th></tr><tr><td>d1</td>
<td>公众号 · 2026-10-15 · https://mp.weixin.qq.com/s/original</td></tr></table>
<h2>风险</h2><p>“雅阁后座靠背偏直。”</p><script>hidden-js</script></body></html>'''


@pytest.fixture(autouse=True)
def configured_export(monkeypatch):
    monkeypatch.setenv('YUQING_BILLING_SERVICE_TOKEN', 'offline-export-auth')


def export_client():
    return TestClient(report.app, headers={'X-Internal-Token': 'offline-export-auth'})


def test_saved_export_keeps_all_sections_and_source_table_without_model_calls(monkeypatch):
    def forbidden(*_args, **_kwargs):
        raise AssertionError('downloading a saved report cannot call the model')
    monkeypatch.setattr(report, 'build_client', forbidden)
    response = export_client().post('/export', json={'html_content': HTML, 'format': 'docx'})
    assert response.status_code == 200
    document = Document(io.BytesIO(response.content))
    text = '\n'.join(p.text for p in document.paragraphs)
    cells = '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert '盘古舆情·第一版' in text and '摘要保持“原文”。\n第二行' in text
    assert '“雅阁后座靠背偏直。”' in text
    assert 'd1' in cells and '2026-10-15' in cells and 'https://mp.weixin.qq.com/s/original' in cells
    assert all(value not in text + cells for value in ['hidden-css', 'hidden-title', 'hidden-js'])


@pytest.mark.parametrize('body', [
    {'html_content': '', 'format': 'docx'},
    {'html_content': '  ', 'format': 'docx'},
    {'html_content': '<p>saved</p>', 'format': 'pdf'},
])
def test_saved_export_rejects_empty_or_unimplemented_exports(body):
    assert export_client().post('/export', json=body).status_code == 422


def test_saved_export_requires_internal_auth_and_does_not_fetch_links():
    client = TestClient(report.app)
    assert client.post('/export', json={'html_content': HTML, 'format': 'docx'}).status_code == 403
    response = client.post('/export', headers={'X-Internal-Token': 'offline-export-auth'},
        json={'html_content': '<p>保持文本<img src="http://192.168.1.1/private"></p>', 'format': 'docx'})
    assert response.status_code == 200
    assert '保持文本' in '\n'.join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)


def test_saved_export_without_configured_auth_fails_closed(monkeypatch):
    monkeypatch.delenv('YUQING_BILLING_SERVICE_TOKEN')
    response = export_client().post('/export', json={'html_content': HTML, 'format': 'docx'})
    assert response.status_code == 503


def test_saved_export_preserves_quote_whitespace_in_table_cells():
    response = export_client().post('/export', json={'html_content': '<table><tr><td>  后排\n实测  </td></tr></table>', 'format': 'docx'})
    assert response.status_code == 200
    assert Document(io.BytesIO(response.content)).tables[0].cell(0, 0).text == '  后排\n实测  '
