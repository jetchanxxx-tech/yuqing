"""Exported reports retain provenance even for documents not quoted by the model."""
import io
from html.parser import HTMLParser

from docx import Document
import pytest
from fastapi.testclient import TestClient

from engines.report_engine import main as report
from engines.tests.test_material_coverage import ReadingProvider


DOCS = [
    {"id": "uncited-source-alpha-01", "title": "系列价格讨论", "content": "新品价格反馈。", "source_type": "toutiao", "source_name": "今日头条",
     "url": "https://www.toutiao.com/article/123/", "final_url": "https://www.toutiao.com/article/456/", "published_at": "2026-10-14T16:00:00Z",
     "source_published_at": "2026年10月15日", "date_sources": {"bocha": "2026-10-15", "page": "2026年10月15日"},
     "content_kind": "summary", "fetch_status": "body_unavailable", "fetch_reason": "blocked"},
    {"id": "uncited-source-beta-02", "title": "续航体验", "content": "本次续航实测。", "source_type": "wechat", "source_name": "公众号",
     "url": "https://mp.weixin.qq.com/s/source-test", "final_url": "https://mp.weixin.qq.com/s/source-test", "published_at": "2026-10-13T10:00:00Z",
     "date_sources": {"page": "2026-10-13T10:00:00Z"}, "content_kind": "body"},
]


class VisibleText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.text = []

    def handle_data(self, data):
        self.text.append(data)


@pytest.mark.parametrize("format", ["html", "docx"])
def test_export_keeps_every_uncited_document_id_date_and_original_final_url(monkeypatch, format):
    provider = ReadingProvider()
    monkeypatch.setattr(report, "build_client", lambda *_args, **_kwargs: provider)
    response = TestClient(report.app).post("/generate", json={"api_key": "offline-key", "documents": DOCS, "format": format})
    assert response.status_code == 200
    if format == "html":
        parsed = VisibleText()
        parsed.feed(response.json()["content"])
        text = "\n".join(parsed.text)
    else:
        document = Document(io.BytesIO(response.content))
        text = "\n".join([paragraph.text for paragraph in document.paragraphs] +
                         [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    for source in DOCS:
        assert source["id"] in text
        assert source["published_at"] in text
        assert source["url"] in text and source["final_url"] in text
        assert all(value in text for value in source["date_sources"].values())
    assert "2026年10月15日" in text
    assert "摘要" in text and "正文未取得" in text
    assert len(provider.calls) == 1, "source row rendering must not purchase another model response"


def test_new_provenance_fields_are_escaped_and_cannot_add_executable_html(monkeypatch):
    provider = ReadingProvider()
    monkeypatch.setattr(report, "build_client", lambda *_args, **_kwargs: provider)
    hostile = {**DOCS[0], "source_published_at": "<img src=x onerror=alert(1)>",
               "date_sources": {"page": "<script>alert('date')</script>"}, "final_url": "javascript:alert('url')"}
    response = TestClient(report.app).post("/generate", json={"api_key": "offline-key", "documents": [hostile]})
    assert response.status_code == 200
    html = response.json()["content"]
    assert "&lt;img" in html and "&lt;script&gt;" in html
    assert "<script>" not in html and "<img src=x" not in html and 'href="javascript:' not in html
