"""Render the immutable published report as Word without model or network calls."""
import io
from html.parser import HTMLParser

from docx import Document
from docx.shared import Pt


class SavedReportParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.document = Document()
        self.document.styles['Normal'].font.size = Pt(10)
        self.parts = []
        self.heading = 0
        self.ignored = 0
        self.table = None
        self.row = None
        self.cell = None

    def flush(self):
        if self.table is not None:
            return
        text = ''.join(self.parts).strip()
        self.parts = []
        if text:
            if self.heading:
                self.document.add_heading(text, level=self.heading)
            else:
                self.document.add_paragraph(text)

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style', 'head'):
            self.ignored += 1
        if self.ignored:
            return
        if tag == 'table':
            self.flush()
            self.table = []
        elif tag == 'tr' and self.table is not None:
            self.row = []
        elif tag in ('td', 'th') and self.row is not None:
            self.cell = []
        elif tag == 'br':
            self.handle_data('\n')
        elif tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6') and self.table is None:
            self.flush()
            self.heading = int(tag[1])
        elif tag in ('p', 'div', 'section', 'li') and self.table is None:
            self.flush()

    def handle_endtag(self, tag):
        if tag in ('script', 'style', 'head') and self.ignored:
            self.ignored -= 1
            return
        if self.ignored:
            return
        if tag in ('td', 'th') and self.cell is not None:
            self.row.append(''.join(self.cell).strip())
            self.cell = None
        elif tag == 'tr' and self.row is not None:
            self.table.append(self.row)
            self.row = None
        elif tag == 'table' and self.table is not None:
            rows = self.table
            self.table = None
            if rows and max(map(len, rows)):
                output = self.document.add_table(rows=0, cols=max(map(len, rows)))
                output.style = 'Light Grid Accent 1'
                for row in rows:
                    cells = output.add_row().cells
                    for index, text in enumerate(row):
                        cells[index].text = text
        elif tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
            self.flush()
            self.heading = 0
        elif tag in ('p', 'div', 'section', 'li'):
            if self.cell is not None:
                self.cell.append('\n')
            else:
                self.flush()

    def handle_data(self, data):
        if self.ignored:
            return
        if self.cell is not None:
            self.cell.append(data)
        elif self.table is None:
            self.parts.append(data)


def render_saved_docx(content):
    parser = SavedReportParser()
    parser.feed(content)
    parser.close()
    parser.flush()
    buffer = io.BytesIO()
    parser.document.save(buffer)
    return buffer.getvalue()
