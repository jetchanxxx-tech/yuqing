#!/usr/bin/env python3
"""Author editable promotional PPTX files and offline HTML previews on a hosted runner.

Run only in .github/workflows/build-release.yml. The single content source is
presentations.json. No network, supplier API, model, notification or payment call
is made by this generator. python-pptx 1.0.2 and a CJK font are required for PPTX.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import os
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DATA = Path(__file__).with_name("presentations.json")
WIDTH = 13.333333
HEIGHT = 7.5
COLORS = {
    "accent": "FF2442",
    "ink": "222224",
    "body": "515156",
    "muted": "66666C",
    "surface": "F5F5F5",
    "warm": "FAFAFA",
    "line": "E3E3E5",
    "white": "FFFFFF",
}


def escaped(value: str) -> str:
    return html.escape(str(value), quote=True)


def paragraph(value: str) -> str:
    return escaped(value).replace("\n", "<br>")


def content_html(slide: dict) -> str:
    """Keep previews and editable slides tied to the same authored data."""
    kind = slide["kind"]
    items = slide["items"]
    if kind == "report":
        return (
            '<div class="report-layout"><article class="sample-report">'
            '<span class="sample-badge">虚构报告结构展示</span>'
            f'<h3>{escaped(slide["sample_title"])}</h3>'
            '<div class="report-rule"></div>'
            f'<p>{paragraph(slide["sample_summary"])}</p>'
            '<div class="report-tags"><span>材料范围</span><span>研判摘要</span>'
            '<span>原始来源</span></div></article><div class="report-parts">'
            + "".join(
                f'<article><span class="item-number">0{i + 1}</span>'
                f'<h3>{escaped(item["title"])}</h3>'
                f'<p>{paragraph(item["body"])}</p></article>'
                for i, item in enumerate(items)
            )
            + "</div></div>"
        )
    if kind == "pricing":
        inner = "".join(
            '<article class="price-card">'
            f'<h3>{escaped(item["title"])}</h3>'
            f'<p class="price">{escaped(item["price"])}</p>'
            f'<p class="credits">{escaped(item["credits"])}</p>'
            f'<p>{paragraph(item["body"])}</p></article>'
            for item in items
        )
        return f'<div class="pricing-grid">{inner}</div><p class="strip">{escaped(slide["strip"])}</p>'
    inner = "".join(
        '<article>'
        f'<span class="item-number">0{i + 1}</span>'
        f'<h3>{escaped(item["title"])}</h3>'
        f'<p>{paragraph(item["body"])}</p></article>'
        for i, item in enumerate(items)
    )
    result = f'<div class="items-grid {kind}-grid">{inner}</div>'
    if slide.get("strip"):
        result += f'<p class="strip">{escaped(slide["strip"])}</p>'
    if slide.get("link"):
        result += (
            f'<a class="product-link" href="{escaped(slide["link"])}" '
            'target="_blank" rel="noopener noreferrer">'
            f'{escaped(slide["link_label"])} <span aria-hidden="true">↗</span></a>'
        )
    return result


HTML_CSS = """
:root{--accent:#ff2442;--ink:#222224;--body:#515156;--muted:#66666c;--line:#e3e3e5;
--canvas:#f5f5f5;--font:"Noto Sans CJK SC","Microsoft YaHei","PingFang SC",sans-serif;
--title-size:clamp(1.5rem,3.4vw,3.25rem);--h3-size:clamp(1rem,1.7vw,1.65rem);
--body-size:clamp(.85rem,1.3vw,1.25rem);--small-size:clamp(.76rem,.95vw,1rem);
--slide-padding:clamp(1rem,4vw,4rem);--content-gap:clamp(.7rem,2.5vh,2rem)}
*{box-sizing:border-box}html,body{margin:0;height:100%;overflow-x:hidden}
html{scroll-snap-type:y mandatory;scroll-behavior:smooth}body{font-family:var(--font);color:var(--ink);background:#fff}
a{color:inherit;text-decoration-thickness:.08em;text-underline-offset:.2em}
button,a{-webkit-tap-highlight-color:transparent}button:focus-visible,a:focus-visible{outline:3px solid #222;outline-offset:4px}
.skip{position:fixed;left:1rem;top:-8rem;background:#fff;padding:.8rem;z-index:8}.skip:focus{top:1rem}
.slide{width:100vw;height:100vh;height:100dvh;overflow:hidden;scroll-snap-align:start;
display:flex;flex-direction:column;position:relative;padding:var(--slide-padding);padding-bottom:clamp(5.8rem,13vh,8rem)}
.slide-content{flex:1;display:flex;flex-direction:column;justify-content:center;max-height:100%;min-height:0;gap:var(--content-gap)}
.topline{display:flex;justify-content:space-between;gap:1rem;font-size:var(--small-size);align-items:center;padding-bottom:clamp(.7rem,2.2vh,1.8rem)}
.brand{font-weight:600;display:flex;align-items:center;gap:.65em}.brand-dot{width:.65em;height:.65em;border-radius:50%;background:var(--accent)}
.audience{color:var(--muted)}.eyebrow{margin:0;color:var(--body);font-size:var(--small-size)}
h1,h2,h3,p{margin:0}h1,h2{font-weight:600;font-size:var(--title-size);line-height:1.3;letter-spacing:0}
.hero h1{font-size:clamp(2.25rem,4.6vw,4.5rem)}.subtitle{font-size:var(--body-size);line-height:1.6;color:var(--body);max-width:70rem}
.items-grid,.pricing-grid{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:clamp(.7rem,1.6vw,1.5rem);margin-top:clamp(.2rem,1vh,1rem)}
.items-grid article,.price-card{background:var(--canvas);padding:clamp(.9rem,2vw,2rem);border-radius:clamp(.7rem,1vw,1rem);min-width:0}
.item-number{display:block;color:var(--muted);font-size:var(--small-size);font-variant-numeric:tabular-nums;margin-bottom:clamp(.5rem,1.7vh,1.2rem)}
h3{font-size:var(--h3-size);font-weight:600;line-height:1.4;margin-bottom:clamp(.45rem,1.3vh,.9rem)}
.items-grid p,.price-card p,.report-layout p{font-size:var(--body-size);color:var(--body);line-height:1.6}
.workflow-grid{grid-template-columns:repeat(4,minmax(0,1fr))}.workflow-grid article{position:relative}
.workflow-grid article+article:before{content:"→";position:absolute;left:calc(clamp(.7rem,1.6vw,1.5rem) * -.85);top:50%;color:var(--muted);font-size:var(--h3-size)}
.metrics-grid{grid-template-columns:repeat(2,minmax(0,1fr))}.metrics-grid article{display:grid;grid-template-columns:auto 1fr;gap:0 .8rem}
.metrics-grid .item-number{grid-row:span 2;margin:.25em 0 0}.metrics-grid h3{margin-bottom:.3em}.metrics-grid p{grid-column:2}
.boundary-grid{grid-template-columns:repeat(2,minmax(0,1fr))}.boundary-grid article:last-child{background:#fafafa;border:1px solid var(--line)}
.price-card h3{color:var(--body)}.price-card .price{font-size:clamp(2rem,3.5vw,3.6rem);color:var(--ink);font-weight:600;line-height:1.4;font-variant-numeric:tabular-nums}
.price-card .credits{font-size:var(--h3-size);color:var(--ink);margin:.3em 0 .65em}.strip{font-size:var(--small-size);line-height:1.6;color:var(--body);padding:clamp(.5rem,1vh,1rem) 0}
.report-layout{display:grid;grid-template-columns:1.1fr 1fr;gap:clamp(1rem,3vw,3rem);align-items:stretch}
.sample-report{background:var(--canvas);padding:clamp(1.2rem,2.5vw,2.8rem);border-radius:1rem}.sample-badge{display:block;font-size:var(--small-size);color:var(--muted);margin-bottom:1rem}
.report-rule{height:1px;background:#d5d5d8;margin:1rem 0}.report-tags{display:flex;flex-wrap:wrap;gap:.5rem;margin-top:1.3rem}.report-tags span{font-size:var(--small-size);background:#fff;padding:.4rem .7rem;border-radius:2rem;color:var(--body)}
.report-parts{display:flex;flex-direction:column;justify-content:space-between;gap:1rem}.report-parts article{display:grid;grid-template-columns:auto 1fr;gap:0 1rem;padding:.2rem 0}
.report-parts .item-number{grid-row:span 2;margin:.3rem 0 0}.report-parts h3{margin:0 0 .3rem}.report-parts p{grid-column:2}
.product-link{align-self:flex-start;background:var(--accent);color:white;padding:clamp(.7rem,1.5vh,1rem) clamp(1rem,2vw,2rem);border-radius:4rem;font-size:var(--body-size);font-weight:500;text-decoration:none}
.product-link:hover{background:#e6203a}.footnote{font-size:var(--small-size);line-height:1.5;color:var(--muted)}
.slide footer{position:absolute;left:var(--slide-padding);right:var(--slide-padding);bottom:clamp(4.6rem,10vh,6.3rem);border-top:1px solid var(--line);padding-top:.65rem}
.controls{position:fixed;inset:auto 0 0;display:flex;justify-content:space-between;align-items:center;gap:.8rem;
background:rgba(255,255,255,.96);border-top:1px solid var(--line);padding:.7rem clamp(1rem,4vw,4rem);z-index:5;font-size:var(--small-size)}
.control-group{display:flex;align-items:center;gap:.7rem}.controls button{font:inherit;border:1px solid #d6d6da;background:#fff;border-radius:4rem;min-height:44px;padding:.55rem 1rem;cursor:pointer;color:var(--ink)}
.controls button:disabled{cursor:default;color:#777;background:#f5f5f5}.counter{font-variant-numeric:tabular-nums;color:var(--body);min-width:4em;text-align:center}
.progress{position:fixed;bottom:0;left:0;height:3px;background:var(--accent);z-index:6;transition:width .18s ease}.keyboard-hint{color:var(--muted)}
.hero .items-grid article{background:#fff;border:1px solid var(--line)}.hero .items-grid .item-number{display:none}
@media(max-height:700px){:root{--slide-padding:clamp(.75rem,3vw,2rem);--content-gap:clamp(.4rem,1.5vh,1rem);--title-size:clamp(1.4rem,3vw,2.6rem)}.topline{padding-bottom:.7rem}.items-grid article,.price-card{padding:clamp(.65rem,1.8vw,1.5rem)}.item-number{margin-bottom:.5rem}}
@media(max-height:600px){:root{--body-size:clamp(.8rem,1.2vw,1.1rem);--slide-padding:clamp(.65rem,2.5vw,1.5rem);--content-gap:.55rem}.hero h1{font-size:clamp(1.8rem,4vw,3rem)}.items-grid p,.report-layout p,.price-card p{line-height:1.45}.keyboard-hint{display:none}.slide{padding-bottom:6rem}}
@media(max-height:500px){:root{--title-size:clamp(1.2rem,2.4vw,1.8rem);--h3-size:clamp(.9rem,1.6vw,1.2rem);--body-size:clamp(.75rem,1.25vw,.95rem);--small-size:clamp(.7rem,1vw,.85rem);--content-gap:.35rem;--slide-padding:.6rem}.topline{padding-bottom:.3rem}.hero h1{font-size:clamp(1.55rem,3.5vw,2rem)}.item-number{margin-bottom:.2rem}.items-grid article,.price-card{padding:.55rem}.slide footer{bottom:4.3rem}.slide{padding-bottom:6rem}.report-tags{margin-top:.6rem}.sample-report{padding:.8rem}.sample-badge{margin-bottom:.6rem}.report-rule{margin:.5rem 0}.product-link{padding:.5rem 1rem}.price-card .price{font-size:1.8rem}}
@media(max-width:600px) and (min-height:501px){:root{--title-size:clamp(1.5rem,6.6vw,2rem);--body-size:clamp(.85rem,3.5vw,1rem);--h3-size:clamp(1rem,4.4vw,1.25rem);--slide-padding:1rem;--content-gap:.65rem}.hero h1{font-size:clamp(1.95rem,8vw,2.6rem)}.items-grid,.pricing-grid,.workflow-grid,.boundary-grid{grid-template-columns:1fr;gap:.5rem}.items-grid article{padding:.65rem .85rem;display:grid;grid-template-columns:auto 1fr;gap:0 .65rem}.items-grid .item-number{grid-row:span 2;margin:.2rem 0 0}.items-grid h3{margin-bottom:.2rem}.items-grid p{grid-column:2;font-size:.83rem;line-height:1.45}.workflow-grid article+article:before{display:none}.hero .items-grid article{display:block}.hero .items-grid h3{display:inline;margin-right:.8rem}.hero .items-grid p{display:inline;font-size:.86rem}.price-card{display:grid;grid-template-columns:1fr auto;gap:.1rem .7rem;padding:.6rem .85rem}.price-card h3{margin:0}.price-card .price{grid-column:2;grid-row:span 2;font-size:1.8rem;line-height:1.4}.price-card .credits{font-size:1rem;margin:0}.price-card p:last-child{font-size:.8rem;grid-column:span 2}.report-layout{grid-template-columns:1fr;gap:.8rem}.sample-report{padding:.9rem}.sample-badge{margin-bottom:.5rem}.report-tags{display:none}.report-rule{margin:.5rem 0}.report-parts{gap:.6rem}.report-parts p{font-size:.83rem;line-height:1.4}.strip{font-size:.76rem;padding:0}.controls .keyboard-hint,.audience{display:none}.controls{padding:.6rem 1rem}.controls button{padding:.5rem .8rem}.slide footer{bottom:4.6rem}.slide{padding-bottom:7rem}.boundary-grid p{font-size:.8rem}.metrics-grid{grid-template-columns:1fr}.product-link{font-size:.87rem;padding:.65rem .85rem}}
@media(max-width:600px) and (min-height:501px){.workflow-grid,.metrics-grid{grid-template-columns:repeat(2,minmax(0,1fr))}.workflow-grid article,.metrics-grid article{display:block;padding:.65rem .75rem}.workflow-grid .item-number,.metrics-grid .item-number{margin:0 0 .3rem}.workflow-grid p,.metrics-grid p{font-size:.85rem}.workflow-grid p br,.metrics-grid p br{display:none}.report .slide-content{gap:.5rem}.report .subtitle{font-size:.85rem}.report-parts article{gap:0 .65rem}.report-parts h3{font-size:1rem}.report .sample-report p{font-size:.83rem}.report .sample-report h3{font-size:1rem}.report .report-layout{gap:.5rem}.report .sample-report{padding:.75rem}.report .report-parts{gap:.45rem}}
@media(max-height:500px){.topline{display:none}.items-grid .item-number{display:none}.items-grid article,.price-card{padding:.55rem .65rem}.hero .items-grid h3{margin-bottom:.2rem}.report-tags{display:none}.report-parts{gap:.35rem}.report-parts .item-number{display:none}.report-parts article{display:block}.sample-badge{margin-bottom:.3rem}.sample-report{padding:.65rem}.metrics-grid article{display:block}.metrics-grid p{grid-column:auto}.hero .slide-content{gap:.3rem}.hero .subtitle{font-size:.75rem}.hero h1{font-size:1.5rem}.product-link{font-size:.75rem;padding:.5rem .75rem}}
@media(prefers-reduced-motion:reduce){html{scroll-behavior:auto}*,*:before,*:after{animation-duration:.01ms!important;transition-duration:.01ms!important}}
@media print{html,body{height:auto;overflow:visible;scroll-snap-type:none}.slide{width:100%;height:7.5in;page-break-after:always;break-after:page;padding:.4in}.slide footer{bottom:.3in}.controls,.progress,.skip{display:none}.slide-content{justify-content:center}.items-grid,.pricing-grid{grid-template-columns:repeat(3,1fr)}.workflow-grid{grid-template-columns:repeat(4,1fr)}.report-layout{grid-template-columns:1.1fr 1fr}h1,h2{font-size:28pt}.hero h1{font-size:38pt}@page{size:13.333in 7.5in;margin:0}}
"""


HTML_JS = """
class PresentationController {
  constructor() {
    this.slides = Array.from(document.querySelectorAll('.slide'));
    this.index = 0;
    this.touchY = null;
    this.touchX = null;
    this.lastWheel = 0;
    this.prev = document.querySelector('[data-prev]');
    this.next = document.querySelector('[data-next]');
    this.counter = document.querySelector('[data-page-count]');
    this.progress = document.querySelector('.progress');
    this.reduced = window.matchMedia('(prefers-reduced-motion: reduce)');
    this.prev.addEventListener('click', () => this.go(this.index - 1));
    this.next.addEventListener('click', () => this.go(this.index + 1));
    document.addEventListener('keydown', event => {
      if (event.altKey || event.ctrlKey || event.metaKey || /INPUT|TEXTAREA|SELECT/.test(event.target.tagName)) return;
      if (['ArrowRight', 'ArrowDown', 'PageDown'].includes(event.key)) {event.preventDefault(); this.go(this.index + 1);}
      if (['ArrowLeft', 'ArrowUp', 'PageUp'].includes(event.key)) {event.preventDefault(); this.go(this.index - 1);}
      if (event.key === 'Home') {event.preventDefault(); this.go(0);}
      if (event.key === 'End') {event.preventDefault(); this.go(this.slides.length - 1);}
    });
    document.addEventListener('wheel', event => {
      if (Math.abs(event.deltaY) < 12 || event.ctrlKey || event.target.closest('.controls')) return;
      event.preventDefault();
      const now = Date.now();
      if (now - this.lastWheel > 650) {this.lastWheel = now; this.go(this.index + Math.sign(event.deltaY));}
    }, {passive: false});
    document.addEventListener('touchstart', event => {
      if (event.target.closest('.controls') || event.touches.length !== 1) return;
      this.touchY = event.touches[0].clientY;
      this.touchX = event.touches[0].clientX;
    }, {passive: true});
    document.addEventListener('touchend', event => {
      if (this.touchY === null || event.changedTouches.length !== 1) return;
      const dy = this.touchY - event.changedTouches[0].clientY;
      const dx = this.touchX - event.changedTouches[0].clientX;
      this.touchY = null;
      this.touchX = null;
      if (Math.max(Math.abs(dx), Math.abs(dy)) > 60) this.go(this.index + Math.sign(Math.abs(dx) > Math.abs(dy) ? dx : dy));
    }, {passive: true});
    this.observer = new IntersectionObserver(entries => {
      for (const entry of entries) {
        if (entry.isIntersecting && entry.intersectionRatio > .6) this.update(this.slides.indexOf(entry.target));
      }
    }, {threshold: [.6, .85]});
    this.slides.forEach(slide => this.observer.observe(slide));
    window.addEventListener('hashchange', () => this.fromHash());
    this.fromHash();
    this.update(this.index);
  }
  fromHash() {
    const match = location.hash.match(/^#slide-(\\d+)$/);
    if (match) this.go(Number(match[1]) - 1);
  }
  go(index) {
    const safe = Math.max(0, Math.min(this.slides.length - 1, index));
    this.update(safe);
    this.slides[safe].scrollIntoView({behavior: this.reduced.matches ? 'instant' : 'smooth', block: 'start'});
  }
  update(index) {
    this.index = index;
    this.prev.disabled = index === 0;
    this.next.disabled = index === this.slides.length - 1;
    this.counter.textContent = `${index + 1} / ${this.slides.length}`;
    this.progress.style.width = `${(index + 1) / this.slides.length * 100}%`;
    this.slides.forEach((slide, i) => slide.setAttribute('aria-current', i === index ? 'page' : 'false'));
    history.replaceState(null, '', `#slide-${index + 1}`);
  }
}
new PresentationController();
"""


def write_html(data: dict, deck: dict, output: Path) -> None:
    sections = []
    for i, slide in enumerate(deck["slides"]):
        tag = "h1" if i == 0 else "h2"
        sections.append(
            f'<section class="slide {escaped(slide["kind"])}" id="slide-{i + 1}" '
            f'aria-labelledby="heading-{i + 1}" data-slide-index="{i + 1}">'
            '<div class="topline"><span class="brand"><span class="brand-dot" aria-hidden="true"></span>'
            f'{escaped(data["brand"])}</span><span class="audience">'
            f'{escaped(deck["audience"])} · {escaped(data["version"])}</span></div>'
            '<div class="slide-content">'
            f'<p class="eyebrow">{escaped(slide["eyebrow"])}</p>'
            f'<{tag} id="heading-{i + 1}">{paragraph(slide["title"])}</{tag}>'
            f'<p class="subtitle">{paragraph(slide["subtitle"])}</p>'
            f'{content_html(slide)}</div><footer><p class="footnote">'
            f'{escaped(slide["footnote"])}</p></footer></section>'
        )
    doc = (
        '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        f'<title>{escaped(deck["title"])}</title>'
        f'<meta name="description" content="{escaped(deck["description"])}">'
        f'<style>{HTML_CSS}</style></head><body><a class="skip" href="#slide-1">跳到演示内容</a>'
        f'<main aria-label="{escaped(deck["title"])}">{"".join(sections)}</main>'
        '<nav class="controls" aria-label="演示翻页"><div class="control-group">'
        '<a href="index.html">材料目录</a><span class="keyboard-hint">方向键翻页 · Home / End</span></div>'
        '<div class="control-group"><button type="button" data-prev aria-label="上一页">上一页</button>'
        '<span class="counter" data-page-count aria-live="polite">1 / '
        f'{len(deck["slides"])}</span><button type="button" data-next aria-label="下一页">下一页</button>'
        f'</div></nav><div class="progress" aria-hidden="true"></div><script>{HTML_JS}</script></body></html>\n'
    )
    output.write_text(doc, encoding="utf-8")


def write_pptx(data: dict, deck: dict, output: Path) -> None:
    # Imports intentionally stay inside the hosted-runner authoring path.
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(WIDTH)
    prs.slide_height = Inches(HEIGHT)
    props = prs.core_properties
    props.title = deck["title"]
    props.subject = deck["description"]
    props.author = data["brand"]
    props.keywords = "盘古舆情,公开信息,AI研判,HTML报告,0.2.5-beta"
    props.comments = "Editable text and vector shapes. Generated on GitHub-hosted runner."
    font = data["font"]

    def rgb(color: str):
        return RGBColor.from_string(COLORS.get(color, color))

    def text(slide, value, x, y, w, h, size=20, color="ink", bold=False, align=None, link=None):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for idx, line in enumerate(str(value).split("\n")):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.space_before = Pt(0)
            p.space_after = Pt(0)
            p.line_spacing = 1.25
            if align is not None:
                p.alignment = align
            run = p.add_run()
            run.text = line
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
            rpr = run._r.get_or_add_rPr()
            rpr.set("lang", "zh-CN")
            for node_name in ("a:ea", "a:cs"):
                node = OxmlElement(node_name)
                node.set("typeface", font)
                rpr.append(node)
            if link:
                run.hyperlink.address = link
        return shape

    def box(slide, x, y, w, h, fill="surface", border=None, radius=True):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if border:
            shape.line.color.rgb = rgb(border)
            shape.line.width = Pt(0.75)
        else:
            shape.line.fill.background()
        if radius:
            shape.adjustments[0] = 0.08
        return shape

    def rule(slide, x1, y1, x2, y2, color="line", weight=1):
        shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        shape.line.color.rgb = rgb(color)
        shape.line.width = Pt(weight)
        return shape

    def circle(slide, x, y, d, fill="accent"):
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.fill.background()
        return shape

    def frame(slide, spec, number):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb("white")
        circle(slide, 0.65, 0.37, 0.1)
        text(slide, data["brand"], 0.86, 0.25, 3.0, 0.35, 16, bold=True)
        text(slide, f'{deck["audience"]} · {data["version"]}', 8.9, 0.26, 3.75, 0.35, 12, "muted", align=PP_ALIGN.RIGHT)
        rule(slide, 0.65, 0.76, 12.68, 0.76)
        rule(slide, 0.65, 6.89, 12.68, 6.89)
        text(slide, spec["footnote"], 0.65, 7.03, 11.45, 0.3, 10, "muted")
        text(slide, f'{number:02d} / {len(deck["slides"]):02d}', 12.0, 7.03, 0.67, 0.3, 10, "muted", align=PP_ALIGN.RIGHT)
        if spec["kind"] != "hero":
            text(slide, spec["eyebrow"], 0.65, 1.05, 12.0, 0.32, 13, "muted")
            text(slide, spec["title"], 0.65, 1.56, 12.0, 0.65, 29, bold=True)
            text(slide, spec["subtitle"], 0.65, 2.36, 12.0, 0.58, 18, "body")

    def draw_hero(slide, spec):
        text(slide, spec["eyebrow"], 0.65, 1.12, 12.0, 0.35, 16, "body")
        text(slide, spec["title"], 0.65, 1.82, 12.0, 1.7, 42, bold=True)
        rule(slide, 0.68, 3.95, 1.5, 3.95, "accent", 3.5)
        text(slide, spec["subtitle"], 0.65, 4.24, 11.8, 0.85, 19, "body")
        for i, item in enumerate(spec["items"]):
            x = 0.65 + i * 4.09
            box(slide, x, 5.45, 3.83, 1.05, fill="white", border="line")
            text(slide, item["title"], x + 0.24, 5.62, 3.3, 0.34, 19, bold=True)
            text(slide, item["body"], x + 0.24, 6.03, 3.3, 0.27, 14, "body")

    def draw_cards(slide, spec, architecture=False):
        items = spec["items"]
        wide = spec["kind"] == "boundary"
        count = len(items)
        gap = 0.24
        w = (12.03 - gap * (count - 1)) / count
        h = 2.87 if not spec.get("strip") else 2.63
        y = 3.22
        for i, item in enumerate(items):
            x = 0.65 + i * (w + gap)
            box(slide, x, y, w, h, fill="warm" if wide and i == 1 else "surface", border="line" if wide and i == 1 else None)
            text(slide, f"0{i + 1}", x + 0.27, y + 0.25, w - 0.54, 0.3, 13, "muted")
            text(slide, item["title"], x + 0.27, y + 0.93, w - 0.54, 0.54, 21, bold=True)
            text(slide, item["body"], x + 0.27, y + 1.61, w - 0.54, 1.12, 18 if not wide else 19, "body")
        if spec.get("strip"):
            text(slide, spec["strip"], 0.65, 6.19, 12.0, 0.37, 16, "body")

    def draw_workflow(slide, spec):
        for i, item in enumerate(spec["items"]):
            x = 0.65 + i * 3.1
            box(slide, x, 3.3, 2.72, 2.77)
            text(slide, f"0{i + 1}", x + 0.24, 3.54, 2.23, 0.5, 25, "muted")
            text(slide, item["title"], x + 0.24, 4.32, 2.23, 0.5, 21, bold=True)
            text(slide, item["body"], x + 0.24, 5.05, 2.23, 0.8, 16.5, "body")
            if i < 3:
                text(slide, "→", x + 2.78, 4.45, 0.33, 0.5, 21, "muted")

    def draw_report(slide, spec):
        box(slide, 0.65, 3.23, 6.05, 3.2)
        text(slide, "虚构报告结构展示", 0.98, 3.5, 5.36, 0.3, 13, "muted")
        text(slide, spec["sample_title"], 0.98, 4.02, 5.36, 0.5, 23, bold=True)
        rule(slide, 0.98, 4.72, 6.3, 4.72)
        text(slide, spec["sample_summary"], 0.98, 4.98, 5.36, 0.85, 17, "body")
        for i, label in enumerate(("材料范围", "研判摘要", "原始来源")):
            x = 1.0 + i * 1.68
            box(slide, x, 6.0, 1.48, 0.3, "white")
            text(slide, label, x, 6.025, 1.48, 0.25, 10, "body", align=PP_ALIGN.CENTER)
        for i, item in enumerate(spec["items"]):
            y = 3.39 + i * 1.07
            text(slide, f"0{i + 1}", 7.2, y + 0.03, 0.4, 0.31, 13, "muted")
            text(slide, item["title"], 7.87, y, 4.8, 0.38, 20, bold=True)
            text(slide, item["body"], 7.87, y + 0.52, 4.8, 0.43, 16, "body")

    def draw_pricing(slide, spec):
        for i, item in enumerate(spec["items"]):
            x = 0.65 + i * 4.09
            box(slide, x, 3.21, 3.83, 2.74)
            text(slide, item["title"], x + 0.27, 3.5, 3.29, 0.4, 18, "body")
            text(slide, item["price"], x + 0.27, 4.01, 3.29, 0.6, 34, bold=True)
            text(slide, item["credits"], x + 0.27, 4.89, 3.29, 0.42, 22, bold=True)
            text(slide, item["body"], x + 0.27, 5.48, 3.29, 0.3, 14, "body")
        text(slide, spec["strip"], 0.65, 6.22, 12.0, 0.4, 16, "body")

    def draw_cta(slide, spec):
        has_link = bool(spec.get("link"))
        y = 3.26
        for i, item in enumerate(spec["items"]):
            x = 0.65 + i * 4.09
            circle(slide, x + 0.03, y + 0.07, 0.12, "ink")
            if i < 2:
                rule(slide, x + 0.25, y + 0.13, x + 3.85, y + 0.13)
            text(slide, item["title"], x, y + 0.52, 3.83, 0.48, 22, bold=True)
            text(slide, item["body"], x, y + 1.24, 3.83, 1.0, 18, "body")
        if has_link:
            box(slide, 0.65, 5.75, 6.15, 0.57, "accent")
            text(slide, spec["link_label"] + " ↗", 0.88, 5.88, 5.71, 0.29, 14, "white", bold=True, link=spec["link"])
            text(slide, spec["link"], 7.14, 5.93, 5.51, 0.36, 12, "body", link=spec["link"])

    def draw_metrics(slide, spec):
        for i, item in enumerate(spec["items"]):
            x = 0.65 + (i % 2) * 6.18
            y = 3.2 + (i // 2) * 1.62
            box(slide, x, y, 5.86, 1.42)
            text(slide, f"0{i + 1}", x + 0.23, y + 0.23, 0.39, 0.3, 13, "muted")
            text(slide, item["title"], x + 0.88, y + 0.21, 4.69, 0.37, 19, bold=True)
            text(slide, item["body"], x + 0.88, y + 0.73, 4.69, 0.6, 15.5, "body")

    drawers = {
        "hero": draw_hero,
        "cards": draw_cards,
        "workflow": draw_workflow,
        "report": draw_report,
        "pricing": draw_pricing,
        "cta": draw_cta,
        "boundary": draw_cards,
        "milestones": draw_cta,
        "architecture": draw_cards,
        "metrics": draw_metrics,
    }
    for idx, spec in enumerate(deck["slides"]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame(slide, spec, idx + 1)
        drawers[spec["kind"]](slide, spec)
        slide.notes_slide.notes_text_frame.text = (
            f'{deck["title"]} — 第 {idx + 1} 页\n'
            f'更新日期：{data["updated"]}；版本：{data["version"]}\n\n'
            f'{spec["notes"]}\n\n依据／限制：\n'
            + "\n".join(spec["sources"])
            + "\n\n字体：Noto Sans CJK SC。如 Windows 未安装此字体，可在 PowerPoint 替换为 Microsoft YaHei。所有文字和图形可编辑。"
        )
    prs.save(output)


def validate_content(data: dict) -> None:
    expected_counts = {"customer": 6, "partner": 10, "investor": 10}
    if {deck["id"] for deck in data["decks"]} != set(expected_counts):
        raise ValueError("All three requested audience decks are required")
    for deck in data["decks"]:
        if len(deck["slides"]) != expected_counts[deck["id"]]:
            raise ValueError(f'Unexpected slide count for {deck["id"]}')
        for number, slide in enumerate(deck["slides"], 1):
            required = ("kind", "title", "subtitle", "items", "footnote", "notes", "sources")
            if any(not slide.get(key) for key in required):
                raise ValueError(f'Incomplete slide: {deck["id"]} {number}')
            if not 2 <= len(slide["items"]) <= 4:
                raise ValueError(f'Slide density outside brief: {deck["id"]} {number}')


def main() -> None:
    if os.environ.get("GITHUB_ACTIONS") != "true" or os.environ.get("RUNNER_ENVIRONMENT") != "github-hosted":
        raise SystemExit("Presentation generation requires the GitHub hosted build-release workflow")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "docs/marketing")
    parser.add_argument("--data", type=Path, default=DEFAULT_DATA)
    args = parser.parse_args()
    data = json.loads(args.data.read_text(encoding="utf-8"))
    validate_content(data)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    manifest = {"brand": data["brand"], "version": data["version"], "updated": data["updated"], "decks": []}
    for deck in data["decks"]:
        base = f'pangu-{deck["id"]}-0.2.5'
        pptx_path = args.output_dir / f"{base}.pptx"
        html_path = args.output_dir / f"{base}.html"
        write_pptx(data, deck, pptx_path)
        write_html(data, deck, html_path)
        manifest["decks"].append({
            "id": deck["id"], "audience": deck["audience"], "title": deck["title"],
            "slides": len(deck["slides"]), "pptx": pptx_path.name, "html": html_path.name,
            "sha256": hashlib.sha256(pptx_path.read_bytes()).hexdigest(),
            "sources": [{"slide": i + 1, "title": slide["title"], "sources": slide["sources"]}
                        for i, slide in enumerate(deck["slides"])],
        })
        print(f'{deck["audience"]}: {len(deck["slides"])} slides -> {pptx_path}')
    (args.output_dir / "presentation-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
